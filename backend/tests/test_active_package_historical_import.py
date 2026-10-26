
from __future__ import annotations

import base64
import io
from contextlib import contextmanager
from datetime import UTC, datetime, timedelta
from uuid import uuid4
from zoneinfo import ZoneInfo

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.agents.v2.semantic_context import build_semantic_context
from app.core.config import settings
from app.integrations.clinic.tia_database_laser import TiaDatabaseLaserClinicAdapter
from app.models.appointment import Appointment
from app.models.branch import Branch
from app.models.clinic_inventory import ClinicLaserDevice, ServiceDevicePrice
from app.models.doctor import Doctor
from app.models.patient import Patient
from app.models.patient_package import PackageUsage, PatientPackage
from app.models.payment_transaction import PaymentTransaction
from app.models.pulse_billing import PatientPulsePack, PulsePackOffer
from app.models.service import Service
from app.models.service_package_offer import ServicePackageOffer
from app.models.staff import Staff
from app.models.user import User
from app.models.workspace import Workspace
from app.schemas.crm import normalize_patient_identity_phone
from app.schemas.historical_import import HistoricalImportDocument
from app.services import historical_import as history
from app.services import inventory as inventory_service
from app.services import package_offers as package_offer_service
from app.services import patient_packages as package_service
from app.services import pulse_billing as pulse_service
from app.services.agent_v2.package_booking_policy import resolve_booking_package
from app.services.agent_v2.planner import PlanStep, ReadRequest
from app.services.agent_v2.read_executor import ReadExecutionContext, execute_step_reads

CAIRO = ZoneInfo("Africa/Cairo")


@contextmanager
def _db_session():
    engine = create_engine(settings.database_url, pool_pre_ping=True)
    connection = engine.connect()
    outer = connection.begin()
    db = Session(
        bind=connection,
        expire_on_commit=False,
        join_transaction_mode="create_savepoint",
    )
    try:
        yield db
    finally:
        db.close()
        if outer.is_active:
            outer.rollback()
        connection.close()
        engine.dispose()


def _seed_workspace(db: Session, *, services: tuple[str, ...] = ("Hydrafacial",)):
    suffix = uuid4().hex[:12]
    user = User(
        email=f"active-package-{suffix}@example.com",
        full_name="Import Test",
        is_active=True,
    )
    workspace = Workspace(
        name=f"Active Package {suffix}",
        slug=f"active-package-{suffix}",
        timezone="Africa/Cairo",
        is_active=True,
        is_demo=True,
    )
    db.add_all([user, workspace])
    db.flush()

    branch = Branch(
        workspace_id=workspace.id,
        name="Main",
        code="MAIN",
        phone=None,
        email=None,
        address_line1=None,
        address_line2=None,
        city="Cairo",
        state=None,
        country_code="EG",
        timezone="Africa/Cairo",
        is_active=True,
    )
    db.add(branch)
    db.flush()
    workspace.primary_branch_id = branch.id

    service_rows: dict[str, Service] = {}
    for index, name in enumerate(services, start=1):
        service = Service(
            workspace_id=workspace.id,
            name=name,
            slug=f"{name.casefold().replace(' ', '-')}-{index}-{suffix}",
            category="Skin",
            operational_category="dermatology",
            description=None,
            duration_minutes=60,
            buffer_before_minutes=0,
            buffer_after_minutes=0,
            price_minor=180_000,
            currency="EGP",
            requires_medical_review=False,
            requires_laser_device=False,
            is_active=True,
        )
        db.add(service)
        db.flush()
        service_rows[name] = service

    staff = Staff(
        workspace_id=workspace.id,
        user_id=None,
        first_name="Dr.",
        last_name="Test",
        email=None,
        phone=None,
        job_title="doctor",
        is_active=True,
    )
    db.add(staff)
    db.flush()
    doctor = Doctor(
        workspace_id=workspace.id,
        staff_id=staff.id,
        doctor_type="regular",
        specialization=None,
        license_number=None,
        bio=None,
        booking_enabled=True,
        is_active=True,
    )
    db.add(doctor)
    db.flush()
    return workspace, user, branch, service_rows, doctor


def _configure_laser_service(
    db: Session,
    *,
    workspace: Workspace,
    service: Service,
) -> dict[str, ServiceDevicePrice]:
    service.operational_category = "laser"
    service.requires_laser_device = True
    db.add_all(
        [
            ClinicLaserDevice(
                workspace_id=workspace.id,
                device_key="prime_lase",
                name="Prime Lase",
                is_active=True,
            ),
            ClinicLaserDevice(
                workspace_id=workspace.id,
                device_key="candela_gentle",
                name="Candela Gentle",
                is_active=True,
            ),
        ]
    )
    db.flush()
    rows = {
        "Prime Lase": ServiceDevicePrice(
            workspace_id=workspace.id,
            service_id=service.id,
            device_key="prime_lase",
            device_name="Prime Lase",
            price_minor=220_000,
            duration_minutes=60,
            currency="EGP",
            is_active=True,
        ),
        "Candela Gentle": ServiceDevicePrice(
            workspace_id=workspace.id,
            service_id=service.id,
            device_key="candela_gentle",
            device_name="Candela Gentle",
            price_minor=260_000,
            duration_minutes=60,
            currency="EGP",
            is_active=True,
        ),
    }
    db.add_all(rows.values())
    db.flush()
    return rows


def _active_package_document(
    rows: list[list[object]],
    *,
    include_package_price: bool = True,
) -> HistoricalImportDocument:
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.append(["test"])
    sheet = wb.create_sheet("active_packages")
    headers = [
        "full_name",
        "phone",
        "service_name",
        "sessions_total",
        "sessions_remaining",
        "amount_paid",
        "purchased_at",
    ]
    if include_package_price:
        headers.append("package_price")
    sheet.append(headers)
    for row in rows:
        sheet.append(row[: len(headers)])
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return HistoricalImportDocument(
        name="active_packages.xlsx",
        format="xlsx",
        content_base64=base64.b64encode(stream.getvalue()).decode("ascii"),
    )


def _arabic_active_package_document(
    rows: list[list[object]],
) -> HistoricalImportDocument:
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.append(["اختبار"])
    sheet = wb.create_sheet("active_packages")
    sheet.append(
        [
            "اسم العميل",
            "رقم الموبايل",
            "اسم الخدمة",
            "اسم الجهاز",
            "عدد الجلسات الكلي",
            "عدد الجلسات المتبقي",
            "المبلغ المدفوع",
            "تاريخ الشراء",
            "سعر الباقة",
            "النبضات المتبقية",
        ]
    )
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return HistoricalImportDocument(
        name="active_packages_ar.xlsx",
        format="xlsx",
        content_base64=base64.b64encode(stream.getvalue()).decode("ascii"),
    )


def _arabic_active_package_document_with_id(
    rows: list[list[object]],
    *,
    id_header: str = "ID",
) -> HistoricalImportDocument:
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.append(["اختبار"])
    sheet = wb.create_sheet("active_packages")
    sheet.append(
        [
            id_header,
            "اسم العميل",
            "رقم الموبايل",
            "اسم الخدمة",
            "اسم الجهاز",
            "عدد الجلسات الكلي",
            "عدد الجلسات المتبقي",
            "المبلغ المدفوع",
            "تاريخ الشراء",
            "سعر الباقة",
            "النبضات المتبقية",
        ]
    )
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return HistoricalImportDocument(
        name="active_packages_id.xlsx",
        format="xlsx",
        content_base64=base64.b64encode(stream.getvalue()).decode("ascii"),
    )


def _id_active_row(
    customer_id: object,
    name: str,
    phone: object,
    *,
    service: object = None,
    device: object = None,
    sessions_total: object = None,
    sessions_remaining: object = None,
    amount_paid: object = None,
    purchased_at: object = None,
    package_price: object = None,
    pulses_remaining: object = None,
) -> list[object]:
    return [
        customer_id,
        name,
        phone,
        service,
        device,
        sessions_total,
        sessions_remaining,
        amount_paid,
        purchased_at,
        package_price,
        pulses_remaining,
    ]


def _unified_document(rows: list[list[object]]) -> HistoricalImportDocument:
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.append(["test"])
    sheet = wb.create_sheet("import_data")
    sheet.append(
        [
            "اسم العميل",
            "رقم الموبايل",
            "اسم الباقة / الخدمة",
            "الجلسات المتبقية",
            "النبضات المتبقية",
            "الجهاز",
        ]
    )
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return HistoricalImportDocument(
        name="unified_history.xlsx",
        format="xlsx",
        content_base64=base64.b64encode(stream.getvalue()).decode("ascii"),
    )


def _legacy_document() -> HistoricalImportDocument:
    wb = Workbook()
    patients = wb.active
    patients.title = "patients"
    patients.append(["full_name", "phone"])
    patients.append(["Legacy Patient", "01019999999"])
    packages = wb.create_sheet("packages")
    packages.append(
        [
            "patient_phone",
            "service_name",
            "sessions_total",
            "sessions_remaining",
            "price",
            "purchased_at",
            "status",
        ]
    )
    packages.append(
        ["01019999999", "Hydrafacial", 6, 2, 3000, "2026-06-01", "active"]
    )
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return HistoricalImportDocument(
        name="legacy.xlsx",
        format="xlsx",
        content_base64=base64.b64encode(stream.getvalue()).decode("ascii"),
    )


def _preview_and_apply(
    db: Session,
    *,
    workspace: Workspace,
    user: User,
    document: HistoricalImportDocument,
    mode: str = "append",
):
    preview = history.preview_historical_import(
        db,
        workspace=workspace,
        user_id=user.id,
        documents=[document],
        mode=mode,
    )
    batch = history.get_historical_import_batch(
        db,
        workspace_id=workspace.id,
        batch_id=preview.batch.batch_id,
    )
    assert batch is not None
    summary = history.apply_historical_import(
        db,
        workspace=workspace,
        batch=batch,
    )
    return preview, summary


@pytest.mark.parametrize(
    ("value", "expected_display", "expected_normalized"),
    [
        ("01012345678", "01012345678", "+201012345678"),
        (1012345678, "01012345678", "+201012345678"),
        ("1012345678", "01012345678", "+201012345678"),
        ("+201012345678", "01012345678", "+201012345678"),
        ("201012345678", "01012345678", "+201012345678"),
        ("00201012345678", "01012345678", "+201012345678"),
    ],
)
def test_historical_import_t18_active_package_egypt_phone_normalization_repairs_excel_leading_zero(
    value,
    expected_display: str,
    expected_normalized: str,
) -> None:
    display, normalized = history._normalize_active_package_egypt_phone(value)
    assert display == expected_display
    assert normalized == expected_normalized


@pytest.mark.parametrize(
    "value",
    [
        "12",
        "abc",
        "01312345678",
        "+971501234567",
        "0212345678",
    ],
)
def test_active_package_phone_rejects_non_egyptian_mobile_numbers(value) -> None:
    with pytest.raises(ValueError):
        history._normalize_active_package_egypt_phone(value)


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("15/2/2026", "2026-02-15"),
        ("15-02-2026", "2026-02-15"),
        ("15.02.2026", "2026-02-15"),
        ("2026-02-15", "2026-02-15"),
        ("2026/02/15", "2026-02-15"),
        ("2026.02.15", "2026-02-15"),
        ("2026-02-15T09:30:00", "2026-02-15"),
        (datetime(2026, 2, 15, 9, 30), "2026-02-15"),
    ],
)
def test_active_package_purchase_date_accepts_excel_and_clear_egyptian_formats(
    value,
    expected: str,
) -> None:
    parsed = history._parse_active_package_purchase_date(value)
    assert parsed is not None
    assert parsed.isoformat() == expected


@pytest.mark.parametrize("value", ["02/15/2026", "not-a-date", 46000, 2.5, True, "31/2"])
def test_active_package_purchase_date_falls_back_for_unusable_values(value) -> None:
    parsed = history._parse_active_package_purchase_date(
        value,
        reference_date=datetime(2026, 10, 15).date(),
    )
    assert parsed is not None
    assert parsed.isoformat() == "2026-10-15"


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, "2026-10-15"),
        ("", "2026-10-15"),
        ("2/3", "2026-03-02"),
        ("20/10", "2025-10-20"),
        ("15/10", "2026-10-15"),
        ("16/10", "2025-10-16"),
        ("29/2", "2024-02-29"),
        ("2", "2026-02-01"),
        (2, "2026-02-01"),
        ("12", "2025-12-01"),
        (chr(0x0662), "2026-02-01"),
        (f"{chr(0x0662)}/{chr(0x0663)}", "2026-03-02"),
    ],
)
def test_active_package_purchase_date_accepts_blank_day_month_and_month_only(
    value,
    expected: str,
) -> None:
    parsed = history._parse_active_package_purchase_date(
        value,
        reference_date=datetime(2026, 10, 15).date(),
    )
    assert parsed is not None
    assert parsed.isoformat() == expected


def test_historical_import_t1_template_adds_optional_customer_id_first() -> None:
    laser_service = "ليزر إزالة الشعر - جسم كامل سيدات"
    payload = history.build_historical_import_template(
        service_names=["Hydrafacial", laser_service],
        device_names=["Prime Lase", "Candela Gentle", "DEKA Again"],
    )
    workbook = load_workbook(io.BytesIO(payload))
    try:
        assert workbook.sheetnames == ["README", "active_packages"]
        readme = workbook["README"]
        sheet = workbook["active_packages"]
        expected_headers = [
            "ID",
            "اسم العميل",
            "رقم الموبايل",
            "اسم الخدمة",
            "اسم الجهاز",
            "عدد الجلسات الكلي",
            "عدد الجلسات المتبقي",
            "المبلغ المدفوع",
            "تاريخ الشراء",
            "سعر الباقة",
            "النبضات المتبقية",
        ]
        assert readme.sheet_view.rightToLeft is True
        assert readme["A1"].value == "دليل استيراد الباقات النشطة"
        readme_text = "\n".join(
            str(cell.value)
            for row in readme.iter_rows()
            for cell in row
            if cell.value is not None
        )
        assert "اسم العميل ورقم الموبايل فقط" in readme_text
        assert "ID اختياري" in readme_text
        assert "أكثر من ID مختلف" in readme_text
        assert "أكثر من باقة مختلفة" in readme_text
        assert "ترك المبلغ المدفوع فارغًا" in readme_text
        assert "سعر الباقة كمبلغ مدفوع" in readme_text
        assert [readme.cell(6, column).value for column in range(1, 12)] == expected_headers
        assert [sheet.cell(1, column).value for column in range(1, 12)] == expected_headers
        assert [sheet.cell(2, column).value for column in range(1, 12)] == [None] * 11
        assert "ActivePackagesTable" in sheet.tables
        assert sheet.tables["ActivePackagesTable"].ref == "A1:K2"
        assert sheet["A2"].number_format == "@"
        assert sheet["C2"].number_format == "@"
        assert sheet["I2"].number_format == "dd/mm/yyyy"
        assert sheet["K2"].number_format == "0"
        validations = {
            str(validation.sqref): validation
            for validation in sheet.data_validations.dataValidation
        }
        assert set(validations) == {"D2:D5000", "E2:E5000"}
        assert "ServicesTable" in validations["D2:D5000"].formula1
        assert "DevicesTable" in validations["E2:E5000"].formula1
        assert validations["D2:D5000"].allow_blank is True
        assert validations["E2:E5000"].allow_blank is True
    finally:
        workbook.close()


def test_active_packages_template_keeps_empty_lookup_tables_without_helper_text() -> None:
    payload = history.build_historical_import_template()
    workbook = load_workbook(io.BytesIO(payload))
    try:
        readme = workbook["README"]
        sheet = workbook["active_packages"]
        assert "ServicesTable" in readme.tables
        assert readme.tables["ServicesTable"].ref == "A11:A12"
        assert "DevicesTable" in readme.tables
        assert readme.tables["DevicesTable"].ref == "D11:D12"
        assert readme["A12"].value is None
        assert readme["D12"].value is None
        assert readme["F14"].value is None
        validations = {
            str(validation.sqref): validation
            for validation in sheet.data_validations.dataValidation
        }
        assert set(validations) == {"D2:D5000", "E2:E5000"}
        assert "ServicesTable" in validations["D2:D5000"].formula1
        assert "DevicesTable" in validations["E2:E5000"].formula1
    finally:
        workbook.close()


def test_active_holding_normalizes_patient_package_pulses_and_both() -> None:
    service = Service(
        workspace_id=uuid4(),
        name="Hydrafacial",
        slug=f"hydrafacial-{uuid4().hex[:8]}",
        category="Skin",
        operational_category="dermatology",
        description=None,
        duration_minutes=60,
        buffer_before_minutes=0,
        buffer_after_minutes=0,
        price_minor=180_000,
        currency="EGP",
        requires_medical_review=False,
        requires_laser_device=False,
        is_active=True,
    )
    service.id = uuid4()
    device = ClinicLaserDevice(
        workspace_id=service.workspace_id,
        device_key="candela_gentle",
        name="Candela Gentle",
        is_active=True,
    )
    device.id = uuid4()
    services = {service.id: service}
    devices = {"candela gentle": [device]}

    package_fields = {
        "service_name": "Hydrafacial",
        "sessions_total": 6,
        "sessions_remaining": 3,
        "package_price": 3000,
        "purchased_at": "15/02/2026",
    }
    cases = [
        ({"full_name": "Ahmed", "phone": "01011111111"}, "patient_only"),
        ({"full_name": "Mariam", "phone": "01122222222", **package_fields}, "package"),
        ({"full_name": "Sara", "phone": "01233333333", "pulses_remaining": 800, "laser_device_name": "Candela Gentle"}, "pulses"),
        ({"full_name": "Nour", "phone": "01544444444", **package_fields, "pulses_remaining": 500, "laser_device_name": "Candela Gentle"}, "both"),
    ]
    for row, expected_kind in cases:
        payload, code, message = history._normalize_active_holding(
            row,
            services_by_id=services,
            device_prices_by_service_id={},
            devices_by_name=devices,
        )
        assert code is None, message
        assert payload is not None
        assert payload["holding_kind"] == expected_kind

    paid_from_price, code, message = history._normalize_active_holding(
        {"full_name": "Paid", "phone": "01055555555", **package_fields},
        services_by_id=services,
        device_prices_by_service_id={},
        devices_by_name=devices,
    )
    assert code is None, message
    assert paid_from_price is not None
    assert paid_from_price["price_minor"] == 300_000
    assert paid_from_price["historical_payment_amount_minor"] == 300_000


def test_unified_holding_normalizes_patient_package_pulses_and_both() -> None:
    service = Service(
        workspace_id=uuid4(),
        name="Hydrafacial",
        slug=f"hydrafacial-{uuid4().hex[:8]}",
        category="Skin",
        operational_category="dermatology",
        description=None,
        duration_minutes=60,
        buffer_before_minutes=0,
        buffer_after_minutes=0,
        price_minor=180_000,
        currency="EGP",
        requires_medical_review=False,
        requires_laser_device=False,
        is_active=True,
    )
    service.id = uuid4()
    device = ClinicLaserDevice(
        workspace_id=service.workspace_id,
        device_key="candela_gentle",
        name="Candela Gentle",
        is_active=True,
    )
    device.id = uuid4()
    services = {service.id: service}
    devices = {"candela gentle": [device]}

    cases = [
        ({"full_name": "Ahmed", "phone": "01011111111"}, "patient_only"),
        ({"full_name": "Mariam", "phone": "01122222222", "service_name": "Hydrafacial", "sessions_remaining": 3}, "package"),
        ({"full_name": "Sara", "phone": "01233333333", "pulses_remaining": 800, "device_name": "Candela Gentle"}, "pulses"),
        ({"full_name": "Nour", "phone": "01544444444", "service_name": "Hydrafacial", "sessions_remaining": 2, "pulses_remaining": 500, "device_name": "Candela Gentle"}, "both"),
    ]
    for row, expected_kind in cases:
        payload, code, message = history._normalize_unified_holding(
            row,
            services_by_id=services,
            device_prices_by_service_id={},
            devices_by_name=devices,
        )
        assert code is None, message
        assert payload is not None
        assert payload["holding_kind"] == expected_kind


def test_unified_holding_rejects_incomplete_or_negative_balances() -> None:
    service = Service(
        workspace_id=uuid4(), name="Hydrafacial", slug=f"hydra-{uuid4().hex[:8]}",
        category="Skin", operational_category="dermatology", description=None, duration_minutes=60,
        buffer_before_minutes=0, buffer_after_minutes=0, price_minor=100_000, currency="EGP",
        requires_medical_review=False, requires_laser_device=False, is_active=True,
    )
    service.id = uuid4()
    services = {service.id: service}
    invalid_rows = [
        {"full_name": "A", "phone": "01011111111", "service_name": "Hydrafacial"},
        {"full_name": "B", "phone": "01122222222", "service_name": "Hydrafacial", "sessions_remaining": -1},
        {"full_name": "C", "phone": "01233333333", "pulses_remaining": -1, "device_name": "Candela Gentle"},
        {"full_name": "D", "phone": "01544444444", "pulses_remaining": 500},
    ]
    for row in invalid_rows:
        payload, code, _message = history._normalize_unified_holding(
            row, services_by_id=services, device_prices_by_service_id={}, devices_by_name={}
        )
        assert payload is None
        assert code is not None


def test_active_package_arabic_headers_preview_and_apply() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(db)
        document = _arabic_active_package_document(
            [[
                "سارة أحمد",
                "01012345678",
                "Hydrafacial",
                None,
                6,
                3,
                3000,
                "15/02/2026",
                3000,
            ]]
        )

        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=document,
        )
        assert preview.ready_counts == {"holding": 1}
        package = db.scalar(
            select(PatientPackage).where(PatientPackage.workspace_id == workspace.id)
        )
        assert package is not None
        assert package.service_id == services["Hydrafacial"].id
        assert package.sessions_purchased == 6
        assert package.opening_sessions_remaining == 3
        assert package.sale_price_minor == 300_000
        assert package.standalone_session_price_minor_at_purchase == 180_000
        read = package_service.package_read(db, package, include_financials=True)
        assert read.cancellation_consumed_sessions == 3
        assert read.cancellation_default_charge_minor == 540_000




def test_active_package_laser_device_is_required_and_persisted() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(
            db,
            services=("Laser Full Body",),
        )
        service = services["Laser Full Body"]
        devices = _configure_laser_service(
            db,
            workspace=workspace,
            service=service,
        )
        document = _arabic_active_package_document(
            [[
                "سارة أحمد",
                "01012345678",
                service.name,
                "Candela Gentle",
                8,
                3,
                5000,
                "15/02/2026",
                5000,
            ]]
        )

        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=document,
        )

        assert preview.ready_counts == {"holding": 1}
        package = db.scalar(
            select(PatientPackage).where(PatientPackage.workspace_id == workspace.id)
        )
        assert package is not None
        assert package.laser_device_key == "candela_gentle"
        assert package.laser_device_name == "Candela Gentle"
        assert package.standalone_session_price_minor_at_purchase == int(
            devices["Candela Gentle"].price_minor
        )
        read = package_service.package_read(db, package, include_financials=True)
        assert read.cancellation_consumed_sessions == 5
        assert read.cancellation_default_charge_minor == 5 * 260_000


def test_active_package_laser_device_validation_is_service_specific() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(
            db,
            services=("Laser Full Body", "Hydrafacial"),
        )
        laser_service = services["Laser Full Body"]
        _configure_laser_service(db, workspace=workspace, service=laser_service)

        missing_device = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[
                _arabic_active_package_document(
                    [[
                        "سارة أحمد",
                        "01012345678",
                        laser_service.name,
                        None,
                        8,
                        3,
                        5000,
                        "15/02/2026",
                        5000,
                    ]]
                )
            ],
            mode="append",
        )
        assert missing_device.rejected_counts["holding"] == 1
        assert missing_device.issue_groups[0].code == "active_package_laser_device_missing"

        non_laser_device = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[
                _arabic_active_package_document(
                    [[
                        "مريم أحمد",
                        "01112345678",
                        "Hydrafacial",
                        "Prime Lase",
                        6,
                        2,
                        3000,
                        "16/02/2026",
                        3000,
                    ]]
                )
            ],
            mode="append",
        )
        assert non_laser_device.rejected_counts["holding"] == 1
        assert non_laser_device.issue_groups[0].code == "active_package_device_not_allowed"


def test_active_package_all_invalid_rows_return_actionable_preview() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        preview = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[
                _active_package_document(
                    [[
                        "Bad Phone",
                        12,
                        "Hydrafacial",
                        6,
                        3,
                        3000,
                        "15/2/2026",
                        None,
                    ]]
                )
            ],
            mode="append",
        )
        assert preview.can_import is False
        assert preview.ready_counts.get("holding", 0) == 0
        assert preview.rejected_counts["holding"] == 1
        assert len(preview.issue_groups) == 1
        issue = preview.issue_groups[0]
        assert issue.code == "active_package_phone_invalid"
        assert issue.example_rows == [2]
        assert "موبايل مصري" in issue.message


@pytest.mark.usefixtures("monkeypatch")
def test_active_package_happy_path_financials_and_opening_balance(monkeypatch) -> None:
    monkeypatch.setattr(history, "record_activity_event", lambda *args, **kwargs: None)
    monkeypatch.setattr(package_service, "record_activity_event", lambda *args, **kwargs: None)

    with _db_session() as db:
        workspace, user, branch, services, doctor = _seed_workspace(db)
        service = services["Hydrafacial"]
        document = _active_package_document(
            [
                [
                    "سارة أحمد",
                    "01012345678",
                    "Hydrafacial",
                    10,
                    4,
                    6000,
                    "2026-07-15",
                    None,
                ]
            ],
            include_package_price=False,
        )

        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=document,
        )
        assert preview.ready_counts == {"holding": 1}
        assert preview.rejected_counts == {}
        assert preview.projected_counts == {
            "patients_to_create": 1,
            "existing_patients_reused": 0,
            "patient_only_records": 0,
            "packages_to_create": 1,
            "pulse_balances_to_create": 0,
            "patients_with_both": 0,
            "historical_payments_to_create": 1,
        }
        assert preview.total_amount_imported_minor == 600_000

        _display, normalized = normalize_patient_identity_phone("01012345678")
        patient = db.scalar(
            select(Patient).where(
                Patient.workspace_id == workspace.id,
                Patient.phone_normalized == normalized,
            )
        )
        assert patient is not None
        assert f"{patient.first_name} {patient.last_name or ''}".strip() == "سارة أحمد"

        package = db.scalar(
            select(PatientPackage).where(
                PatientPackage.workspace_id == workspace.id,
                PatientPackage.patient_id == patient.id,
            )
        )
        assert package is not None
        assert package.service_id == service.id
        assert package.sessions_purchased == 10
        assert package.opening_sessions_remaining == 4
        assert package.sessions_total_known is True
        assert package.sale_price_minor == 600_000
        assert package.expires_at is None
        assert package.status == "active"
        assert package.source == "integration"
        assert package.purchase_transaction_id is not None
        assert db.scalar(
            select(func.count())
            .select_from(PackageUsage)
            .where(PackageUsage.patient_package_id == package.id)
        ) == 0

        payment = db.get(PaymentTransaction, package.purchase_transaction_id)
        assert payment is not None
        assert payment.patient_id == patient.id
        assert payment.patient_package_id == package.id
        assert payment.transaction_type == "payment"
        assert payment.amount_minor == 600_000
        assert payment.currency == "EGP"
        assert payment.payment_method == "unknown"
        assert payment.source == "integration"
        assert payment.created_at == package.purchased_at

        financial = package_service.package_read(db, package, include_financials=True)
        assert financial.sessions_remaining == 4
        assert financial.amount_paid_minor == 600_000
        assert financial.balance_due_minor == 0
        assert financial.effective_status == "active"

        agent_read = execute_step_reads(
            PlanStep(
                operation_index=0,
                operation_type="package_info",
                disposition="read",
                reads=[
                    ReadRequest(
                        kind="customer_packages",
                        parameters={"service_id": str(service.id)},
                    )
                ],
                response_goal="package_information",
            ),
            ReadExecutionContext(
                db=db,
                workspace=workspace,
                patient=patient,
                now=datetime.now(UTC),
            ),
        )
        agent_packages = agent_read.results[0].payload["packages"]
        assert len(agent_packages) == 1
        assert agent_packages[0]["sessions_remaining"] == 4

        start_at = datetime.now(UTC) + timedelta(days=10)
        booking_resolution = resolve_booking_package(
            db,
            workspace=workspace,
            patient=patient,
            service_id=service.id,
            start_at=start_at,
            device_key=None,
            package_usage="use_existing",
        )
        assert booking_resolution.package_used is True
        assert booking_resolution.package_id == package.id
        end_at = start_at + timedelta(minutes=60)
        package_service.validate_package_for_booking(
            db,
            workspace_id=workspace.id,
            package_id=package.id,
            patient_id=patient.id,
            service_id=service.id,
            appointment_start_at=start_at,
        )
        appointment = Appointment(
            workspace_id=workspace.id,
            patient_id=patient.id,
            branch_id=branch.id,
            doctor_id=doctor.id,
            doctor_assignment_known=True,
            is_quick_booking=False,
            service_id=service.id,
            patient_package_id=None,
            visit_group_id=None,
            lead_id=None,
            created_by_user_id=user.id,
            rescheduled_from_appointment_id=None,
            status="confirmed",
            source="staff",
            start_at=start_at,
            end_at=end_at,
            busy_start_at=start_at,
            busy_end_at=end_at,
            duration_minutes=60,
            price_minor=service.price_minor,
            discount_minor=0,
            currency="EGP",
            laser_device_key=None,
            laser_device_name=None,
            laser_pulses_used=None,
            payment_status="unknown",
            amount_paid_minor=None,
            payment_method="unknown",
            billing_context="standard",
            package_external_id=None,
            customer_note=None,
            cancellation_reason=None,
            idempotency_key=None,
        )
        db.add(appointment)
        db.flush()
        package_service.reserve_package_usage(
            db,
            appointment=appointment,
            package=package,
            actor_user_id=user.id,
        )
        after = package_service.package_read(db, package, include_financials=True)
        assert after.sessions_remaining == 3
        assert after.amount_paid_minor == 600_000
        assert after.balance_due_minor == 0


def test_active_package_partial_payment_patient_reuse_multiple_and_reimport(monkeypatch) -> None:
    monkeypatch.setattr(history, "record_activity_event", lambda *args, **kwargs: None)
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(db)
        service = services["Hydrafacial"]
        display, normalized = normalize_patient_identity_phone("01012345678")
        existing = Patient(
            workspace_id=workspace.id,
            first_name="Existing",
            last_name="CRM",
            phone=display,
            phone_normalized=normalized,
            gender="female",
            preferred_language="ar",
            preferred_branch_id=workspace.primary_branch_id,
            source="referral",
            status="active",
            marketing_consent=False,
        )
        db.add(existing)
        db.flush()

        rows = [
            [
                "Different Excel Name",
                "01012345678",
                "Hydrafacial",
                10,
                4,
                5000,
                "2026-07-15",
                8000,
            ],
            [
                "Different Excel Name",
                "+201012345678",
                "Hydrafacial",
                6,
                2,
                3000,
                "2026-08-15",
                None,
            ],
        ]
        document = _active_package_document(rows)

        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=document,
        )
        assert preview.projected_counts == {
            "patients_to_create": 0,
            "existing_patients_reused": 1,
            "patient_only_records": 0,
            "packages_to_create": 2,
            "pulse_balances_to_create": 0,
            "patients_with_both": 0,
            "historical_payments_to_create": 2,
        }
        assert preview.total_amount_imported_minor == 800_000

        patients = list(
            db.scalars(
                select(Patient).where(
                    Patient.workspace_id == workspace.id,
                    Patient.phone_normalized == normalized,
                )
            ).all()
        )
        assert len(patients) == 1
        assert patients[0].id == existing.id
        assert patients[0].first_name == "Existing"
        assert patients[0].last_name == "CRM"

        packages = list(
            db.scalars(
                select(PatientPackage)
                .where(
                    PatientPackage.workspace_id == workspace.id,
                    PatientPackage.patient_id == existing.id,
                    PatientPackage.service_id == service.id,
                )
                .order_by(PatientPackage.purchased_at)
            ).all()
        )
        assert len(packages) == 2
        payments = list(
            db.scalars(
                select(PaymentTransaction).where(
                    PaymentTransaction.workspace_id == workspace.id,
                    PaymentTransaction.patient_package_id.in_([row.id for row in packages]),
                )
            ).all()
        )
        assert len(payments) == 2

        partial = next(row for row in packages if row.sale_price_minor == 800_000)
        financial = package_service.package_read(db, partial, include_financials=True)
        assert financial.amount_paid_minor == 500_000
        assert financial.balance_due_minor == 300_000

        renamed_rows = [list(row) for row in rows]
        renamed_rows[0][0] = "Another Historical Name"
        renamed_rows[1][0] = "Another Historical Name"
        preview2, _summary2 = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_active_package_document(renamed_rows),
        )
        assert preview2.projected_counts["packages_to_create"] == 0
        assert preview2.projected_counts["historical_payments_to_create"] == 0
        assert preview2.total_amount_imported_minor == 0
        assert db.scalar(
            select(func.count())
            .select_from(PatientPackage)
            .where(PatientPackage.workspace_id == workspace.id)
        ) == 2
        assert db.scalar(
            select(func.count())
            .select_from(PaymentTransaction)
            .where(PaymentTransaction.workspace_id == workspace.id)
        ) == 2

        reordered = _active_package_document(list(reversed(rows)))
        preview3, _summary3 = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=reordered,
        )
        assert preview3.projected_counts["packages_to_create"] == 0
        assert preview3.projected_counts["historical_payments_to_create"] == 0
        assert db.scalar(
            select(func.count())
            .select_from(PatientPackage)
            .where(PatientPackage.workspace_id == workspace.id)
        ) == 2
        assert db.scalar(
            select(func.count())
            .select_from(PaymentTransaction)
            .where(PaymentTransaction.workspace_id == workspace.id)
        ) == 2


def test_active_package_preview_rejects_invalid_and_duplicate_rows_without_writes() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        valid = [
            "Valid Patient",
            "01011111111",
            "Hydrafacial",
            10,
            4,
            6000,
            "2026-07-15",
            None,
        ]
        rows = [
            valid,
            list(valid),
            ["Unknown", "01011111112", "Unknown", 10, 4, 6000, "2026-07-15", None],
            ["Zero Total", "01011111113", "Hydrafacial", 0, 0, 6000, "2026-07-15", None],
            ["Negative Rem", "01011111114", "Hydrafacial", 10, -1, 6000, "2026-07-15", None],
            ["Too Much Rem", "01011111115", "Hydrafacial", 10, 11, 6000, "2026-07-15", None],
            ["Zero Paid", "01011111116", "Hydrafacial", 10, 4, 0, "2026-07-15", None],
            ["Negative Paid", "01011111117", "Hydrafacial", 10, 4, -1, "2026-07-15", None],
            ["Low Price", "01011111118", "Hydrafacial", 10, 4, 6000, "2026-07-15", 5000],
            ["Bad Date", "01011111119", "Hydrafacial", 10, 4, 6000, "not-a-date", None],
            ["Bad Phone", "abc", "Hydrafacial", 10, 4, 6000, "2026-07-15", None],
        ]
        preview = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[_active_package_document(rows)],
            mode="append",
        )
        assert preview.ready_counts["holding"] == 2
        assert preview.rejected_counts["holding"] == len(rows) - 2
        codes = {issue.code for issue in preview.issue_groups}
        assert {
            "duplicate_historical_package_row",
            "active_package_service_unknown",
            "active_package_total_invalid",
            "active_package_remaining_invalid",
            "active_package_remaining_exceeds_total",
            "active_package_amount_paid_invalid",
            "active_package_price_below_paid",
            "active_package_phone_invalid",
        }.issubset(codes)
        assert db.scalar(
            select(func.count())
            .select_from(Patient)
            .where(Patient.workspace_id == workspace.id)
        ) == 0
        assert db.scalar(
            select(func.count())
            .select_from(PatientPackage)
            .where(PatientPackage.workspace_id == workspace.id)
        ) == 0
        assert db.scalar(
            select(func.count())
            .select_from(PaymentTransaction)
            .where(PaymentTransaction.workspace_id == workspace.id)
        ) == 0


def test_legacy_multisheet_workbook_remains_previewable_and_importable(monkeypatch) -> None:
    monkeypatch.setattr(history, "record_activity_event", lambda *args, **kwargs: None)
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_legacy_document(),
        )
        assert preview.batch.schema_version == "tia_history_v1"
        assert preview.ready_counts["patient"] == 1
        assert preview.ready_counts["package"] == 1
        assert db.scalar(
            select(func.count())
            .select_from(Patient)
            .where(Patient.workspace_id == workspace.id)
        ) == 1
        package = db.scalar(
            select(PatientPackage).where(PatientPackage.workspace_id == workspace.id)
        )
        assert package is not None
        assert package.opening_sessions_remaining == 2
        assert package.source == "integration"


def test_device_rename_updates_current_config_but_preserves_historical_snapshots() -> None:
    with _db_session() as db:
        workspace, user, branch, services, doctor = _seed_workspace(
            db,
            services=("Laser Full Body",),
        )
        service = services["Laser Full Body"]
        service.operational_category = "laser"
        service.requires_laser_device = True

        device = ClinicLaserDevice(
            workspace_id=workspace.id,
            device_key="device_deka_again",
            name="DEKA Again",
            is_active=True,
        )
        db.add(device)
        db.flush()

        price = ServiceDevicePrice(
            workspace_id=workspace.id,
            service_id=service.id,
            device_key=device.device_key,
            device_name="DEKA Again",
            price_minor=155_000,
            duration_minutes=50,
            currency="EGP",
            is_active=True,
        )
        package_offer = ServicePackageOffer(
            workspace_id=workspace.id,
            service_id=service.id,
            device_key=device.device_key,
            device_name="DEKA Again",
            sessions_count=6,
            price_minor=800_000,
            currency="EGP",
            is_active=True,
        )
        pulse_offer = PulsePackOffer(
            workspace_id=workspace.id,
            device_key=device.device_key,
            device_name="DEKA Again",
            pulses_count=1000,
            price_minor=120_000,
            currency="EGP",
            is_active=True,
        )
        db.add_all([price, package_offer, pulse_offer])

        display, normalized = normalize_patient_identity_phone("01022222222")
        patient = Patient(
            workspace_id=workspace.id,
            first_name="Historical",
            last_name="Patient",
            phone=display,
            phone_normalized=normalized,
            gender="female",
            preferred_language="ar",
            preferred_branch_id=branch.id,
            source="referral",
            status="active",
            marketing_consent=False,
        )
        db.add(patient)
        db.flush()

        purchased_at = datetime(2026, 1, 5, 10, 0, tzinfo=UTC)
        historical_package = PatientPackage(
            workspace_id=workspace.id,
            patient_id=patient.id,
            service_id=service.id,
            purchase_transaction_id=None,
            package_offer_id=package_offer.id,
            origin_appointment_id=None,
            created_by_user_id=user.id,
            external_id=None,
            name="Historical DEKA package",
            sessions_purchased=6,
            opening_sessions_remaining=4,
            sessions_total_known=True,
            sale_price_minor=800_000,
            standalone_session_price_minor_at_purchase=155_000,
            laser_device_key=device.device_key,
            laser_device_name="DEKA Again",
            currency="EGP",
            purchased_at=purchased_at,
            expires_at=None,
            status="active",
            source="staff",
            idempotency_key=None,
        )
        historical_pulse_pack = PatientPulsePack(
            workspace_id=workspace.id,
            patient_id=patient.id,
            pulse_pack_offer_id=pulse_offer.id,
            origin_appointment_id=None,
            purchase_transaction_id=None,
            created_by_user_id=user.id,
            device_key=device.device_key,
            device_name="DEKA Again",
            pulses_purchased=1000,
            sale_price_minor=120_000,
            standalone_pulse_price_minor_at_purchase=150,
            currency="EGP",
            purchased_at=purchased_at,
            expires_at=None,
            status="active",
            idempotency_key=None,
        )
        appointment_start = datetime(2026, 1, 10, 10, 0, tzinfo=UTC)
        historical_appointment = Appointment(
            workspace_id=workspace.id,
            patient_id=patient.id,
            branch_id=branch.id,
            doctor_id=doctor.id,
            doctor_assignment_known=True,
            is_quick_booking=False,
            service_id=service.id,
            patient_package_id=None,
            visit_group_id=None,
            lead_id=None,
            created_by_user_id=user.id,
            rescheduled_from_appointment_id=None,
            status="completed",
            source="staff",
            start_at=appointment_start,
            end_at=appointment_start + timedelta(minutes=50),
            busy_start_at=appointment_start,
            busy_end_at=appointment_start + timedelta(minutes=50),
            duration_minutes=50,
            price_minor=155_000,
            discount_minor=0,
            currency="EGP",
            laser_device_key=device.device_key,
            laser_device_name="DEKA Again",
            laser_pulses_used=500,
            payment_status="paid",
            amount_paid_minor=155_000,
            payment_method="cash",
            billing_context="standard",
            package_external_id=None,
            customer_note=None,
            cancellation_reason=None,
            idempotency_key=None,
        )
        db.add_all(
            [historical_package, historical_pulse_pack, historical_appointment]
        )
        db.flush()

        device_id = device.id
        package_id = historical_package.id
        pulse_pack_id = historical_pulse_pack.id
        appointment_id = historical_appointment.id
        device_key = device.device_key

        renamed = inventory_service.update_clinic_laser_device(
            db,
            workspace_id=workspace.id,
            device_id=device_id,
            name="DEKA Again Pro",
            is_active=None,
        )
        assert renamed.device_key == device_key
        assert renamed.name == "DEKA Again Pro"

        db.expire_all()

        current_prices = inventory_service.list_laser_device_prices(
            db,
            workspace_id=workspace.id,
        )
        current_price = next(
            row for row in current_prices if row.device_key == device_key
        )
        assert current_price.device_name == "DEKA Again Pro"

        current_package_offer = next(
            row
            for row in package_offer_service.list_package_offers(
                db,
                workspace_id=workspace.id,
                active_only=True,
            )
            if row.device_key == device_key
        )
        assert current_package_offer.device_name == "DEKA Again Pro"

        current_pulse_offer = next(
            row
            for row in pulse_service.list_pulse_pack_offers(
                db,
                workspace_id=workspace.id,
                active_only=True,
            )
            if row.device_key == device_key
        )
        assert current_pulse_offer.device_name == "DEKA Again Pro"

        catalog = TiaDatabaseLaserClinicAdapter(
            db=db,
            workspace=workspace,
        ).build_catalog()
        semantic = build_semantic_context(catalog)
        assert {"ref": "V1", "name": "DEKA Again Pro"} in semantic.model_input[
            "devices"
        ]

        current_device_names = list(
            db.scalars(
                select(ClinicLaserDevice.name).where(
                    ClinicLaserDevice.workspace_id == workspace.id,
                    ClinicLaserDevice.is_active.is_(True),
                )
            )
        )
        template_bytes = history.build_historical_import_template(
            service_names=[service.name],
            device_names=current_device_names,
        )
        workbook = load_workbook(io.BytesIO(template_bytes), data_only=False)
        try:
            workbook_values = {
                str(cell.value)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            }
        finally:
            workbook.close()
        assert "DEKA Again Pro" in workbook_values

        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document(
                [[
                    "Rename Import",
                    "01033333333",
                    service.name,
                    "DEKA Again Pro",
                    6,
                    3,
                    3000,
                    "15/02/2026",
                    3000,
                ]]
            ),
        )
        assert preview.ready_counts == {"holding": 1}
        assert "active_package_laser_device_invalid" not in {
            issue.code for issue in preview.issue_groups
        }
        imported = db.scalar(
            select(PatientPackage)
            .where(
                PatientPackage.workspace_id == workspace.id,
                PatientPackage.source == "integration",
                PatientPackage.laser_device_key == device_key,
            )
            .order_by(PatientPackage.created_at.desc())
        )
        assert imported is not None
        assert imported.laser_device_name == "DEKA Again Pro"

        old_appointment = db.get(Appointment, appointment_id)
        old_package = db.get(PatientPackage, package_id)
        old_pulse_pack = db.get(PatientPulsePack, pulse_pack_id)
        assert old_appointment is not None
        assert old_package is not None
        assert old_pulse_pack is not None
        assert old_appointment.laser_device_name == "DEKA Again"
        assert old_package.laser_device_name == "DEKA Again"
        assert old_pulse_pack.device_name == "DEKA Again"


def test_active_packages_four_shapes_create_usable_packages_pulses_and_financials() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(
            db, services=("Hydrafacial", "Laser Full Body")
        )
        laser_service = services["Laser Full Body"]
        _configure_laser_service(db, workspace=workspace, service=laser_service)
        document = _arabic_active_package_document(
            [
                ["Ahmed Patient", "01011111111", None, None, None, None, None, None, None, None],
                ["Mariam Package", "01122222222", "Hydrafacial", None, 6, 3, None, "15/02/2026", 3000, None],
                ["Sara Pulses", "01233333333", None, "Candela Gentle", None, None, None, None, None, 800],
                ["Nour Both", "01544444444", "Laser Full Body", "Candela Gentle", 8, 2, None, "16/02/2026", 5000, 500],
            ]
        )
        preview, summary = _preview_and_apply(
            db, workspace=workspace, user=user, document=document
        )
        assert preview.ready_counts == {"holding": 4}
        assert preview.rejected_counts == {}
        assert preview.projected_counts["patients_to_create"] == 4
        assert preview.projected_counts["patient_only_records"] == 1
        assert preview.projected_counts["packages_to_create"] == 2
        assert preview.projected_counts["pulse_balances_to_create"] == 2
        assert preview.projected_counts["patients_with_both"] == 1
        assert preview.projected_counts["historical_payments_to_create"] == 2
        assert preview.total_amount_imported_minor == 800_000
        assert summary["imported_counts"]["package"] == 2
        assert summary["imported_counts"]["payment"] == 2
        assert summary["imported_counts"]["pulse_pack"] == 2

        normalized_phones = {
            "+201011111111", "+201122222222", "+201233333333", "+201544444444"
        }
        patients = list(
            db.scalars(
                select(Patient).where(
                    Patient.workspace_id == workspace.id,
                    Patient.phone_normalized.in_(normalized_phones),
                )
            ).all()
        )
        assert len(patients) == 4
        patient_ids = {patient.id for patient in patients}
        packages = list(
            db.scalars(
                select(PatientPackage).where(
                    PatientPackage.workspace_id == workspace.id,
                    PatientPackage.patient_id.in_(patient_ids),
                )
            ).all()
        )
        pulse_packs = list(
            db.scalars(
                select(PatientPulsePack).where(
                    PatientPulsePack.workspace_id == workspace.id,
                    PatientPulsePack.patient_id.in_(patient_ids),
                )
            ).all()
        )
        payments = list(
            db.scalars(
                select(PaymentTransaction).where(
                    PaymentTransaction.workspace_id == workspace.id,
                    PaymentTransaction.patient_id.in_(patient_ids),
                )
            ).all()
        )
        assert len(packages) == 2
        assert sorted(package.opening_sessions_remaining for package in packages) == [2, 3]
        assert all(package.purchase_transaction_id is not None for package in packages)
        assert sorted(package.sale_price_minor for package in packages) == [300_000, 500_000]
        assert len(payments) == 2
        assert sorted(payment.amount_minor for payment in payments) == [300_000, 500_000]
        assert len(pulse_packs) == 2
        assert sorted(pack.pulses_purchased for pack in pulse_packs) == [500, 800]
        assert all(pack.purchase_transaction_id is None for pack in pulse_packs)

        hydra_package = next(package for package in packages if package.service_id == services["Hydrafacial"].id)
        package_read = package_service.package_read(db, hydra_package, include_financials=True)
        assert package_read.sessions_remaining == 3
        assert package_read.amount_paid_minor == 300_000
        assert package_read.balance_due_minor == 0

        pulse_patient = next(patient for patient in patients if patient.phone_normalized == "+201233333333")
        pulse_read = pulse_service.list_patient_pulse_packs(
            db,
            workspace_id=workspace.id,
            patient_id=pulse_patient.id,
        )
        assert len(pulse_read) == 1
        assert pulse_read[0].pulses_remaining == 800


def test_unified_historical_import_four_shapes_create_canonical_records_without_fake_revenue() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(
            db, services=("Hydrafacial", "Laser Full Body")
        )
        laser_service = services["Laser Full Body"]
        _configure_laser_service(db, workspace=workspace, service=laser_service)
        document = _unified_document(
            [
                ["Ahmed Patient", "01011111111", None, None, None, None],
                ["Mariam Package", "01122222222", "Hydrafacial", 3, None, None],
                ["Sara Pulses", "01233333333", None, None, 800, "Candela Gentle"],
                ["Nour Both", "01544444444", "Laser Full Body", 2, 500, "Candela Gentle"],
            ]
        )
        preview, summary = _preview_and_apply(
            db, workspace=workspace, user=user, document=document
        )
        assert preview.ready_counts == {"holding": 4}
        assert preview.rejected_counts == {}
        assert preview.projected_counts["patient_only_records"] == 1
        assert preview.projected_counts["packages_to_create"] == 2
        assert preview.projected_counts["pulse_balances_to_create"] == 2
        assert preview.projected_counts["patients_with_both"] == 1
        assert preview.projected_counts["historical_payments_to_create"] == 0
        assert preview.total_amount_imported_minor == 0
        assert summary["imported_counts"]["package"] == 2
        assert summary["imported_counts"]["pulse_pack"] == 2

        normalized_phones = {
            "+201011111111", "+201122222222", "+201233333333", "+201544444444"
        }
        patients = list(
            db.scalars(
                select(Patient).where(
                    Patient.workspace_id == workspace.id,
                    Patient.phone_normalized.in_(normalized_phones),
                )
            ).all()
        )
        assert len(patients) == 4
        patient_ids = {patient.id for patient in patients}
        packages = list(
            db.scalars(
                select(PatientPackage).where(
                    PatientPackage.workspace_id == workspace.id,
                    PatientPackage.patient_id.in_(patient_ids),
                )
            ).all()
        )
        pulse_packs = list(
            db.scalars(
                select(PatientPulsePack).where(
                    PatientPulsePack.workspace_id == workspace.id,
                    PatientPulsePack.patient_id.in_(patient_ids),
                )
            ).all()
        )
        payments = list(
            db.scalars(
                select(PaymentTransaction).where(
                    PaymentTransaction.workspace_id == workspace.id,
                    PaymentTransaction.patient_id.in_(patient_ids),
                )
            ).all()
        )
        assert len(packages) == 2
        assert all(package.purchase_transaction_id is None for package in packages)
        assert sorted(package.opening_sessions_remaining for package in packages) == [2, 3]
        assert len(pulse_packs) == 2
        assert sorted(pack.pulses_purchased for pack in pulse_packs) == [500, 800]
        assert all(pack.purchase_transaction_id is None for pack in pulse_packs)
        assert payments == []


def test_active_packages_missing_required_header_is_structured_validation_error() -> None:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "active_packages"
    sheet.append([
        "اسم العميل",
        "اسم الخدمة",
        "عدد الجلسات الكلي",
        "عدد الجلسات المتبقي",
        "المبلغ المدفوع",
        "تاريخ الشراء",
        "سعر الباقة",
        "النبضات المتبقية",
    ])
    sheet.append(["سارة أحمد", "Hydrafacial", 6, 3, 3000, "15/02/2026", 3000, None])
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    document = HistoricalImportDocument(
        name="missing-phone.xlsx",
        format="xlsx",
        content_base64=base64.b64encode(stream.getvalue()).decode("ascii"),
    )

    with pytest.raises(history.HistoricalImportValidationError) as exc_info:
        list(history._rows_from_xlsx(document))

    assert exc_info.value.errors == [
        {
            "row": None,
            "field": "phone",
            "message": 'العمود المطلوب "رقم الموبايل" غير موجود.',
        }
    ]


def test_active_packages_multiple_invalid_rows_keep_excel_rows_and_no_canonical_writes() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        document = _arabic_active_package_document(
            [
                [None, "01011111111", None, None, None, None, None, None, None, None],
                ["Bad Phone", "abc", None, None, None, None, None, None, None, None],
                ["Unknown Service", "01122222222", "Not In Linka", None, 6, 3, 3000, "15/02/2026", 3000, None],
                ["Bad Sessions", "01233333333", "Hydrafacial", None, 0, 0, None, "16/02/2026", 3000, None],
            ]
        )

        preview = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[document],
            mode="append",
        )

        assert preview.can_import is False
        assert preview.ready_counts.get("holding", 0) == 0
        assert preview.rejected_counts["holding"] == 4
        rows_by_code = {
            issue.code: issue.example_rows
            for issue in preview.issue_groups
        }
        assert rows_by_code["active_package_name_missing"] == [2]
        assert rows_by_code["active_package_phone_invalid"] == [3]
        assert rows_by_code["active_package_service_unknown"] == [4]
        assert rows_by_code["active_package_total_invalid"] == [5]
        assert db.scalar(
            select(func.count())
            .select_from(Patient)
            .where(Patient.workspace_id == workspace.id)
        ) == 0
        assert db.scalar(
            select(func.count())
            .select_from(PatientPackage)
            .where(PatientPackage.workspace_id == workspace.id)
        ) == 0
        assert db.scalar(
            select(func.count())
            .select_from(PaymentTransaction)
            .where(PaymentTransaction.workspace_id == workspace.id)
        ) == 0
        assert db.scalar(
            select(func.count())
            .select_from(PatientPulsePack)
            .where(PatientPulsePack.workspace_id == workspace.id)
        ) == 0


def test_history_route_structures_file_validation_without_exposing_internals() -> None:
    from app.api.routes.clinic_setup_v2 import _setup_error

    exc = history.HistoricalImportValidationError(
        [
            {"row": 8, "field": "phone", "message": "رقم الهاتف غير صالح."},
            {"row": 12, "field": "service_name", "message": "الخدمة غير موجودة."},
        ]
    )
    response = _setup_error(exc)

    assert response.status_code == 422
    assert response.detail == {
        "code": "import_validation_failed",
        "message": "تعذر استيراد الملف بسبب أخطاء في البيانات.",
        "error_count": 2,
        "errors": exc.errors,
    }


def test_historical_import_t2_customer_only_with_id() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id(
                [_id_active_row("100", "سارة أحمد", "01012345678")]
            ),
        )
        assert preview.ready_counts == {"holding": 1}
        patient = db.scalar(select(Patient).where(Patient.workspace_id == workspace.id))
        assert patient is not None
        assert patient.customer_id == "100"
        assert db.scalar(select(func.count()).select_from(PatientPackage).where(PatientPackage.workspace_id == workspace.id)) == 0
        assert db.scalar(select(func.count()).select_from(PatientPulsePack).where(PatientPulsePack.workspace_id == workspace.id)) == 0


def test_historical_import_t3_customer_only_without_id() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document(
                [["سارة أحمد", "01012345678", None, None, None, None, None, None, None, None]]
            ),
        )
        assert preview.ready_counts == {"holding": 1}
        patient = db.scalar(select(Patient).where(Patient.workspace_id == workspace.id))
        assert patient is not None
        assert patient.customer_id is None


def test_historical_import_t4_t5_customer_id_text_and_aliases() -> None:
    cases = [
        ("ID", "000123", "01011111111"),
        ("id", "CL-0012", "01011111112"),
        ("customer_id", "A-123", "01011111113"),
        ("Customer ID", "CUST/123", "01011111114"),
    ]
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        for header, customer_id, phone in cases:
            _preview, _summary = _preview_and_apply(
                db,
                workspace=workspace,
                user=user,
                document=_arabic_active_package_document_with_id(
                    [_id_active_row(customer_id, f"Patient {customer_id}", phone)],
                    id_header=header,
                ),
            )
        patients = {
            row.phone: row.customer_id
            for row in db.scalars(select(Patient).where(Patient.workspace_id == workspace.id)).all()
        }
        assert patients["01011111111"] == "000123"
        assert patients["01011111112"] == "CL-0012"
        assert patients["01011111113"] == "A-123"
        assert patients["01011111114"] == "CUST/123"


def test_historical_import_t6_same_phone_same_id_one_patient() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        row = _id_active_row("100", "سارة أحمد", "01012345678")
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id([row, list(row)]),
        )
        assert preview.ready_counts["holding"] == 1
        assert db.scalar(select(func.count()).select_from(Patient).where(Patient.workspace_id == workspace.id)) == 1
        patient = db.scalar(select(Patient).where(Patient.workspace_id == workspace.id))
        assert patient is not None and patient.customer_id == "100"


@pytest.mark.parametrize("rows", [
    [
        _id_active_row(None, "سارة أحمد", "01012345678"),
        _id_active_row("100", "سارة أحمد", "01012345678"),
    ],
    [
        _id_active_row("100", "سارة أحمد", "01012345678"),
        _id_active_row(None, "سارة أحمد", "01012345678"),
    ],
])
def test_historical_import_t7_blank_and_supplied_id_resolve_to_supplied(rows) -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id(rows),
        )
        assert preview.ready_counts["holding"] == 1
        patient = db.scalar(select(Patient).where(Patient.workspace_id == workspace.id))
        assert patient is not None
        assert patient.customer_id == "100"
        assert db.scalar(select(func.count()).select_from(Patient).where(Patient.workspace_id == workspace.id)) == 1


def test_historical_import_t8_same_phone_different_ids_preview_conflict() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        preview = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[_arabic_active_package_document_with_id([
                _id_active_row("100", "سارة أحمد", "01012345678"),
                _id_active_row("200", "سارة أحمد", "01012345678"),
            ])],
            mode="append",
        )
        assert preview.can_import is False
        assert preview.ready_counts.get("holding", 0) == 0
        conflict = next(group for group in preview.issue_groups if group.code == "customer_id_conflict_in_file")
        assert conflict.example_rows == [2, 3]
        assert "100" in conflict.message and "200" in conflict.message
        assert db.scalar(select(func.count()).select_from(Patient).where(Patient.workspace_id == workspace.id)) == 0
        batch = history.get_historical_import_batch(db, workspace_id=workspace.id, batch_id=preview.batch.batch_id)
        assert batch is not None
        with pytest.raises(history.HistoricalImportError, match="no ready rows"):
            history.apply_historical_import(db, workspace=workspace, batch=batch)


def _existing_patient(db: Session, workspace: Workspace, *, phone: str, customer_id: str | None) -> Patient:
    display, normalized = normalize_patient_identity_phone(phone)
    patient = Patient(
        workspace_id=workspace.id,
        first_name="Existing",
        last_name="Patient",
        phone=display,
        phone_normalized=normalized,
        customer_id=customer_id,
        gender="female",
        preferred_language="ar",
        preferred_branch_id=workspace.primary_branch_id,
        source="other",
        status="active",
        marketing_consent=False,
    )
    db.add(patient)
    db.flush()
    return patient


def test_historical_import_t9_existing_patient_missing_id_is_filled() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        existing = _existing_patient(db, workspace, phone="01012345678", customer_id=None)
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id([
                _id_active_row("100", "Different Import Name", "01012345678")
            ]),
        )
        assert preview.projected_counts["existing_patients_reused"] == 1
        db.refresh(existing)
        assert existing.customer_id == "100"
        assert db.scalar(select(func.count()).select_from(Patient).where(Patient.workspace_id == workspace.id)) == 1


def test_historical_import_t10_existing_patient_same_id_is_safe_noop() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        existing = _existing_patient(db, workspace, phone="01012345678", customer_id="100")
        _preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id([
                _id_active_row("100", "Different Import Name", "01012345678")
            ]),
        )
        db.refresh(existing)
        assert existing.customer_id == "100"
        assert existing.first_name == "Existing"
        assert existing.last_name == "Patient"
        assert db.scalar(select(func.count()).select_from(Patient).where(Patient.workspace_id == workspace.id)) == 1


def test_historical_import_t11_existing_patient_conflicting_id_is_preview_conflict() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        existing = _existing_patient(db, workspace, phone="01012345678", customer_id="100")
        preview = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[_arabic_active_package_document_with_id([
                _id_active_row("200", "Existing Patient", "01012345678")
            ])],
            mode="append",
        )
        assert preview.can_import is False
        conflict = next(group for group in preview.issue_groups if group.code == "customer_id_conflict_existing_patient")
        assert "Current ID: 100" in conflict.message
        assert "Imported ID: 200" in conflict.message
        db.refresh(existing)
        assert existing.customer_id == "100"
        assert db.scalar(select(func.count()).select_from(Patient).where(Patient.workspace_id == workspace.id)) == 1


def test_historical_import_t12_existing_patient_blank_import_id_preserves_existing() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        existing = _existing_patient(db, workspace, phone="01012345678", customer_id="100")
        _preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id([
                _id_active_row(None, "Existing Patient", "01012345678")
            ]),
        )
        db.refresh(existing)
        assert existing.customer_id == "100"


def test_historical_import_t13_one_phone_two_different_packages() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(db, services=("Hydrafacial", "Peel"))
        rows = [
            _id_active_row("100", "سارة أحمد", "01012345678", service="Hydrafacial", sessions_total=6, sessions_remaining=3, amount_paid=3000, purchased_at="15/02/2026", package_price=3000),
            _id_active_row("100", "سارة أحمد", "01012345678", service="Peel", sessions_total=8, sessions_remaining=5, amount_paid=4000, purchased_at="16/02/2026", package_price=4000),
        ]
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id(rows),
        )
        assert preview.ready_counts["holding"] == 2
        assert db.scalar(select(func.count()).select_from(Patient).where(Patient.workspace_id == workspace.id)) == 1
        packages = list(db.scalars(select(PatientPackage).where(PatientPackage.workspace_id == workspace.id)).all())
        assert len(packages) == 2
        assert {row.service_id for row in packages} == {services["Hydrafacial"].id, services["Peel"].id}


def test_historical_import_t14_exact_duplicate_package_writes_once() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        row = _id_active_row("100", "سارة أحمد", "01012345678", service="Hydrafacial", sessions_total=6, sessions_remaining=3, amount_paid=3000, purchased_at="15/02/2026", package_price=3000)
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id([row, list(row)]),
        )
        assert any(group.code == "duplicate_historical_package_row" for group in preview.issue_groups)
        assert db.scalar(select(func.count()).select_from(PatientPackage).where(PatientPackage.workspace_id == workspace.id)) == 1
        assert db.scalar(select(func.count()).select_from(PaymentTransaction).where(PaymentTransaction.workspace_id == workspace.id)) == 1


def test_historical_import_t15_duplicate_pulse_component_writes_once() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db, services=("Hydrafacial", "Peel"))
        db.add(ClinicLaserDevice(workspace_id=workspace.id, device_key="candela_gentle", name="Candela Gentle", is_active=True))
        db.flush()
        rows = [
            _id_active_row("100", "سارة أحمد", "01012345678", service="Hydrafacial", device="Candela Gentle", sessions_total=6, sessions_remaining=3, amount_paid=3000, purchased_at="15/02/2026", package_price=3000, pulses_remaining=1000),
            _id_active_row("100", "سارة أحمد", "01012345678", service="Peel", device="Candela Gentle", sessions_total=8, sessions_remaining=5, amount_paid=4000, purchased_at="16/02/2026", package_price=4000, pulses_remaining=1000),
        ]
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id(rows),
        )
        assert any(group.code == "duplicate_pulse_pack_holding" for group in preview.issue_groups)
        assert db.scalar(select(func.count()).select_from(PatientPulsePack).where(PatientPulsePack.workspace_id == workspace.id)) == 1


def test_historical_import_t16_same_phone_is_workspace_scoped() -> None:
    with _db_session() as db:
        workspace_a, _user_a, _branch_a, _services_a, _doctor_a = _seed_workspace(db)
        workspace_b, user_b, _branch_b, _services_b, _doctor_b = _seed_workspace(db)
        patient_a = _existing_patient(db, workspace_a, phone="01012345678", customer_id="100")
        _preview, _summary = _preview_and_apply(
            db,
            workspace=workspace_b,
            user=user_b,
            document=_arabic_active_package_document_with_id([
                _id_active_row("200", "Workspace B", "01012345678")
            ]),
        )
        display, normalized = normalize_patient_identity_phone("01012345678")
        del display
        patient_b = db.scalar(select(Patient).where(Patient.workspace_id == workspace_b.id, Patient.phone_normalized == normalized))
        assert patient_b is not None and patient_b.customer_id == "200"
        db.refresh(patient_a)
        assert patient_a.customer_id == "100"
        assert patient_a.id != patient_b.id


def test_historical_import_t17_old_ten_column_template_is_still_accepted() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document([
                ["Legacy Ten Columns", "01012345678", None, None, None, None, None, None, None, None]
            ]),
        )
        assert preview.ready_counts == {"holding": 1}
        patient = db.scalar(select(Patient).where(Patient.workspace_id == workspace.id))
        assert patient is not None and patient.customer_id is None


def test_historical_import_conflicting_names_do_not_duplicate_patient() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(db, services=("Hydrafacial", "Peel"))
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id([
                _id_active_row("100", "سارة أحمد", "01012345678", service="Hydrafacial", sessions_total=6, sessions_remaining=3, amount_paid=3000, purchased_at="15/02/2026", package_price=3000),
                _id_active_row("100", "سارة محمد", "01012345678", service="Peel", sessions_total=8, sessions_remaining=4, amount_paid=4000, purchased_at="16/02/2026", package_price=4000),
            ]),
        )
        assert preview.ready_counts["holding"] == 2
        assert db.scalar(select(func.count()).select_from(Patient).where(Patient.workspace_id == workspace.id)) == 1
        assert db.scalar(select(func.count()).select_from(PatientPackage).where(PatientPackage.workspace_id == workspace.id)) == 2
        patient = db.scalar(select(Patient).where(Patient.workspace_id == workspace.id))
        assert patient is not None
        assert f"{patient.first_name} {patient.last_name or ''}".strip() == "سارة أحمد"
        assert {row.service_id for row in db.scalars(select(PatientPackage).where(PatientPackage.workspace_id == workspace.id)).all()} == {services["Hydrafacial"].id, services["Peel"].id}


def test_historical_import_apply_rechecks_customer_id_after_preview() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        existing = _existing_patient(db, workspace, phone="01012345678", customer_id=None)
        preview = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[_arabic_active_package_document_with_id([
                _id_active_row("100", "Existing Patient", "01012345678", service="Hydrafacial", sessions_total=6, sessions_remaining=3, amount_paid=3000, purchased_at="15/02/2026", package_price=3000)
            ])],
            mode="append",
        )
        existing.customer_id = "999"
        db.commit()
        batch = history.get_historical_import_batch(db, workspace_id=workspace.id, batch_id=preview.batch.batch_id)
        assert batch is not None
        with pytest.raises(history.HistoricalImportConflictError, match="customer ID changed after preview"):
            history.apply_historical_import(db, workspace=workspace, batch=batch)
        db.rollback()
        assert db.scalar(select(func.count()).select_from(PatientPackage).where(PatientPackage.workspace_id == workspace.id)) == 0
        db.refresh(existing)
        assert existing.customer_id == "999"


@pytest.mark.parametrize("header,value", [
    ("ID", "000123"),
    ("id", "CL-0012"),
    ("customer_id", "A-123"),
    ("Customer ID", "CUST/123"),
])
def test_historical_import_customer_id_headers_parse_as_text_without_database(header: str, value: str) -> None:
    document = _arabic_active_package_document_with_id(
        [_id_active_row(value, "Header Test", "01012345678")],
        id_header=header,
    )
    rows = list(history._iter_rows([document]))
    assert len(rows) == 1
    _source_file, sheet, row_number, row = rows[0]
    assert sheet == "active_packages"
    assert row_number == 2
    assert row["customer_id"] == value



def _add_package_offer(
    db: Session,
    *,
    workspace: Workspace,
    service: Service,
    sessions: int,
    price_minor: int,
    device_key: str | None = None,
    active: bool = True,
) -> ServicePackageOffer:
    return package_offer_service.upsert_package_offer(
        db,
        workspace_id=workspace.id,
        service_id=service.id,
        device_key=device_key,
        sessions_count=sessions,
        price_minor=price_minor,
        currency="EGP",
        is_active=active,
    )


def test_historical_package_fallback_f1_missing_financials_uses_exact_offer_and_full_payment() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(db)
        offer = _add_package_offer(
            db,
            workspace=workspace,
            service=services["Hydrafacial"],
            sessions=6,
            price_minor=300_000,
        )
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id([
                _id_active_row(None, "Ahmed", "01011111111", service="Hydrafacial", sessions_total=6, sessions_remaining=2)
            ]),
        )
        assert preview.ready_counts == {"holding": 1}
        package = db.scalar(select(PatientPackage).where(PatientPackage.workspace_id == workspace.id))
        assert package is not None
        assert package.sessions_purchased == 6
        assert package.opening_sessions_remaining == 2
        assert package.sale_price_minor == 300_000
        payment = db.scalar(
            select(PaymentTransaction).where(
                PaymentTransaction.workspace_id == workspace.id,
                PaymentTransaction.patient_package_id == package.id,
                PaymentTransaction.transaction_type == "payment",
            )
        )
        assert payment is not None
        assert payment.amount_minor == 300_000
        assert package.purchase_transaction_id == payment.id
        assert db.scalar(
            select(func.count()).select_from(PaymentTransaction).where(
                PaymentTransaction.workspace_id == workspace.id,
                PaymentTransaction.patient_package_id == package.id,
                PaymentTransaction.transaction_type == "payment",
            )
        ) == 1
        assert offer.price_minor == 300_000


@pytest.mark.parametrize(
    ("case", "package_price", "amount_paid", "expected_price", "expected_paid"),
    [
        ("F2", 2500, None, 250_000, 250_000),
        ("F3", None, 2500, 250_000, 250_000),
        ("F4", 3000, 2500, 300_000, 250_000),
        ("F5", 2500, None, 250_000, 250_000),
    ],
)
def test_historical_package_fallback_f2_to_f5_explicit_financial_precedence(
    case: str,
    package_price: object,
    amount_paid: object,
    expected_price: int,
    expected_paid: int,
) -> None:
    workspace_id = uuid4()
    service = Service(
        workspace_id=workspace_id,
        name="Hydrafacial",
        slug=f"hydra-{uuid4().hex[:8]}",
        category="Skin",
        operational_category="dermatology",
        description=None,
        duration_minutes=60,
        buffer_before_minutes=0,
        buffer_after_minutes=0,
        price_minor=180_000,
        currency="EGP",
        requires_medical_review=False,
        requires_laser_device=False,
        is_active=True,
    )
    service.id = uuid4()
    offer = ServicePackageOffer(
        workspace_id=workspace_id,
        service_id=service.id,
        device_key=None,
        device_name=None,
        sessions_count=6,
        price_minor=300_000,
        currency="EGP",
        is_active=True,
    )
    offer.id = uuid4()
    payload, code, message = history._normalize_active_holding(
        {
            "full_name": case,
            "phone": "01011111111",
            "service_name": "Hydrafacial",
            "sessions_total": 6,
            "sessions_remaining": 2,
            "package_price": package_price,
            "amount_paid": amount_paid,
            "purchased_at": "2026-10-01",
        },
        services_by_id={service.id: service},
        device_prices_by_service_id={},
        devices_by_name={},
        package_offers_by_identity={(service.id, None, 6): offer},
    )
    assert code is None, message
    assert payload is not None
    assert payload["price_minor"] == expected_price
    assert payload["historical_payment_amount_minor"] == expected_paid
    assert payload["financial_fallback_offer_id"] is None
    assert payload["financial_fallback_offer_price_minor"] is None


def test_historical_package_fallback_f6_no_exact_active_offer_blocks_preview_and_writes() -> None:
    with _db_session() as db:
        workspace, user, _branch, _services, _doctor = _seed_workspace(db)
        preview = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[_arabic_active_package_document_with_id([
                _id_active_row(None, "No Offer", "01022222222", service="Hydrafacial", sessions_total=6, sessions_remaining=2)
            ])],
            mode="append",
        )
        assert preview.ready_counts == {}
        assert preview.rejected_counts == {"holding": 1}
        assert {issue.code for issue in preview.issue_groups} == {"active_holding_package_offer_missing"}
        assert db.scalar(select(func.count()).select_from(PatientPackage).where(PatientPackage.workspace_id == workspace.id)) == 0
        assert db.scalar(select(func.count()).select_from(PaymentTransaction).where(PaymentTransaction.workspace_id == workspace.id)) == 0


def test_historical_package_fallback_f7_wrong_session_count_does_not_nearest_match() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(db)
        _add_package_offer(db, workspace=workspace, service=services["Hydrafacial"], sessions=6, price_minor=300_000)
        _add_package_offer(db, workspace=workspace, service=services["Hydrafacial"], sessions=10, price_minor=450_000)
        preview = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[_arabic_active_package_document_with_id([
                _id_active_row(None, "Eight", "01033333333", service="Hydrafacial", sessions_total=8, sessions_remaining=2)
            ])],
            mode="append",
        )
        assert preview.ready_counts == {}
        assert {issue.code for issue in preview.issue_groups} == {"active_holding_package_offer_missing"}


def test_historical_package_fallback_f8_laser_requires_exact_device_offer() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(db, services=("Full Body",))
        service = services["Full Body"]
        _configure_laser_service(db, workspace=workspace, service=service)
        _add_package_offer(db, workspace=workspace, service=service, sessions=8, price_minor=480_000, device_key="prime_lase")
        _add_package_offer(db, workspace=workspace, service=service, sessions=8, price_minor=520_000, device_key="candela_gentle")
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id([
                _id_active_row(None, "Laser", "01044444444", service="Full Body", device="Candela Gentle", sessions_total=8, sessions_remaining=3)
            ]),
        )
        assert preview.ready_counts == {"holding": 1}
        package = db.scalar(select(PatientPackage).where(PatientPackage.workspace_id == workspace.id))
        assert package is not None
        assert package.laser_device_key == "candela_gentle"
        assert package.sale_price_minor == 520_000


def test_historical_package_fallback_f9_inactive_exact_offer_is_not_eligible() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(db)
        _add_package_offer(db, workspace=workspace, service=services["Hydrafacial"], sessions=6, price_minor=300_000, active=False)
        preview = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[_arabic_active_package_document_with_id([
                _id_active_row(None, "Inactive", "01055555555", service="Hydrafacial", sessions_total=6, sessions_remaining=2)
            ])],
            mode="append",
        )
        assert preview.ready_counts == {}
        assert {issue.code for issue in preview.issue_groups} == {"active_holding_package_offer_missing"}


@pytest.mark.parametrize("value", [None, "not-a-date"])
def test_historical_package_fallback_f10_f11_purchase_date_remains_optional(value) -> None:
    reference = datetime(2026, 10, 26).date()
    assert history._parse_active_package_purchase_date(value, reference_date=reference) == reference


def test_historical_package_fallback_f12_offer_resolution_is_workspace_scoped() -> None:
    with _db_session() as db:
        workspace_a, user_a, _branch_a, _services_a, _doctor_a = _seed_workspace(db)
        workspace_b, _user_b, _branch_b, services_b, _doctor_b = _seed_workspace(db)
        _add_package_offer(db, workspace=workspace_b, service=services_b["Hydrafacial"], sessions=6, price_minor=300_000)
        preview = history.preview_historical_import(
            db,
            workspace=workspace_a,
            user_id=user_a.id,
            documents=[_arabic_active_package_document_with_id([
                _id_active_row(None, "Scoped", "01066666666", service="Hydrafacial", sessions_total=6, sessions_remaining=2)
            ])],
            mode="append",
        )
        assert preview.ready_counts == {}
        assert {issue.code for issue in preview.issue_groups} == {"active_holding_package_offer_missing"}
        assert db.scalar(select(func.count()).select_from(PatientPackage).where(PatientPackage.workspace_id == workspace_a.id)) == 0


def test_historical_package_fallback_f13_apply_rechecks_offer_price_before_any_write() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(db)
        offer = _add_package_offer(db, workspace=workspace, service=services["Hydrafacial"], sessions=6, price_minor=300_000)
        preview = history.preview_historical_import(
            db,
            workspace=workspace,
            user_id=user.id,
            documents=[_arabic_active_package_document_with_id([
                _id_active_row(None, "Race", "01077777777", service="Hydrafacial", sessions_total=6, sessions_remaining=2)
            ])],
            mode="append",
        )
        batch = history.get_historical_import_batch(db, workspace_id=workspace.id, batch_id=preview.batch.batch_id)
        assert batch is not None
        offer.price_minor = 350_000
        db.flush()
        with pytest.raises(history.HistoricalImportConflictError, match="changed after preview"):
            history.apply_historical_import(db, workspace=workspace, batch=batch)
        assert db.scalar(select(func.count()).select_from(PatientPackage).where(PatientPackage.workspace_id == workspace.id)) == 0
        assert db.scalar(select(func.count()).select_from(PaymentTransaction).where(PaymentTransaction.workspace_id == workspace.id)) == 0


def test_historical_package_fallback_f14_rerun_does_not_duplicate_package_or_payment() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(db)
        _add_package_offer(db, workspace=workspace, service=services["Hydrafacial"], sessions=6, price_minor=300_000)
        document = _arabic_active_package_document_with_id([
            _id_active_row(None, "Repeat", "01088888888", service="Hydrafacial", sessions_total=6, sessions_remaining=2, purchased_at="2026-10-01")
        ])
        _preview_and_apply(db, workspace=workspace, user=user, document=document)
        _preview_and_apply(db, workspace=workspace, user=user, document=document)
        assert db.scalar(select(func.count()).select_from(PatientPackage).where(PatientPackage.workspace_id == workspace.id)) == 1
        assert db.scalar(select(func.count()).select_from(PaymentTransaction).where(
            PaymentTransaction.workspace_id == workspace.id,
            PaymentTransaction.transaction_type == "payment",
        )) == 1


def test_historical_package_fallback_zero_price_offer_creates_no_fake_payment() -> None:
    with _db_session() as db:
        workspace, user, _branch, services, _doctor = _seed_workspace(db)
        _add_package_offer(db, workspace=workspace, service=services["Hydrafacial"], sessions=6, price_minor=0)
        preview, _summary = _preview_and_apply(
            db,
            workspace=workspace,
            user=user,
            document=_arabic_active_package_document_with_id([
                _id_active_row(None, "Free", "01099999999", service="Hydrafacial", sessions_total=6, sessions_remaining=2)
            ]),
        )
        assert preview.ready_counts == {"holding": 1}
        package = db.scalar(select(PatientPackage).where(PatientPackage.workspace_id == workspace.id))
        assert package is not None and package.sale_price_minor == 0
        assert db.scalar(select(func.count()).select_from(PaymentTransaction).where(PaymentTransaction.workspace_id == workspace.id)) == 0


def test_historical_package_fallback_template_explains_optional_financials_and_date() -> None:
    workbook = load_workbook(io.BytesIO(history.build_historical_import_template(service_names=["Hydrafacial"], device_names=[])))
    try:
        readme_text = "\n".join(
            str(cell.value)
            for row in workbook["README"].iter_rows()
            for cell in row
            if cell.value is not None
        )
        assert "سعر الباقة والمبلغ المدفوع وتاريخ الشراء حقول اختيارية" in readme_text
        assert "إذا تُرك السعر والمبلغ المدفوع فارغين معًا" in readme_text
        assert "المطابقة تمامًا للخدمة والجهاز وعدد الجلسات" in readme_text
    finally:
        workbook.close()
