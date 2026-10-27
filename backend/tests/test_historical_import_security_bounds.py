from __future__ import annotations

import base64
import io
import json

import pytest
from openpyxl import Workbook

from app.core.historical_import_limits import (
    MAX_DOCUMENT_BASE64_CHARS,
    MAX_DOCUMENT_BYTES,
    MAX_DOCUMENT_NAME_CHARS,
    MAX_DOCUMENTS,
    MAX_HISTORICAL_IMPORT_JSON_ENVELOPE_BYTES,
    MAX_HISTORICAL_IMPORT_REQUEST_BYTES,
    MAX_UPLOAD_BASE64_CHARS,
    MAX_UPLOAD_BYTES,
    base64_encoded_length,
    max_base64_chars_for_independent_documents,
)
from app.schemas.historical_import import HistoricalImportDocument, HistoricalImportPreviewRequest
from app.services import historical_import as history


def _constructed_document(*, name: str, format: str, content_base64: str) -> HistoricalImportDocument:
    return HistoricalImportDocument.model_construct(
        name=name,
        format=format,
        content_base64=content_base64,
    )


def test_historical_import_boundary_math_is_exact() -> None:
    assert MAX_DOCUMENT_BYTES == 26_214_400
    assert base64_encoded_length(MAX_DOCUMENT_BYTES) == 34_952_536
    assert MAX_DOCUMENT_BASE64_CHARS == 34_952_536

    assert MAX_UPLOAD_BYTES == 62_914_560
    assert max_base64_chars_for_independent_documents(MAX_UPLOAD_BYTES, MAX_DOCUMENTS) == 83_886_104
    assert MAX_UPLOAD_BASE64_CHARS == 83_886_104


def test_raw_request_cap_is_derived_from_worst_bounded_metadata_envelope() -> None:
    worst_name = "\U0010ffff" * MAX_DOCUMENT_NAME_CHARS
    document = {"name": worst_name, "format": "xlsx", "content_base64": ""}
    payload = {
        "documents": [dict(document) for _ in range(MAX_DOCUMENTS)],
        "mode": "replace_previous_imports",
    }
    envelope = json.dumps(payload, ensure_ascii=True, separators=(",", ":")).encode("utf-8")

    assert len(envelope) == MAX_HISTORICAL_IMPORT_JSON_ENVELOPE_BYTES
    assert MAX_UPLOAD_BASE64_CHARS + len(envelope) == MAX_HISTORICAL_IMPORT_REQUEST_BYTES
    assert MAX_HISTORICAL_IMPORT_REQUEST_BYTES == 83_917_233


def test_schema_uses_canonical_document_and_base64_bounds() -> None:
    document_schema = HistoricalImportDocument.model_json_schema()["properties"]
    preview_schema = HistoricalImportPreviewRequest.model_json_schema()["properties"]

    assert document_schema["name"]["maxLength"] == MAX_DOCUMENT_NAME_CHARS
    assert document_schema["content_base64"]["maxLength"] == MAX_DOCUMENT_BASE64_CHARS
    assert preview_schema["documents"]["maxItems"] == MAX_DOCUMENTS


def test_sec_d1_oversized_encoded_document_never_calls_b64decode(monkeypatch) -> None:
    calls = 0

    def unexpected_decode(*args, **kwargs):
        nonlocal calls
        calls += 1
        raise AssertionError("b64decode must not run")

    monkeypatch.setattr(history, "MAX_DOCUMENT_BASE64_CHARS", 4)
    monkeypatch.setattr(history, "MAX_UPLOAD_BASE64_CHARS", 100)
    monkeypatch.setattr(history.base64, "b64decode", unexpected_decode)
    document = _constructed_document(name="patients.csv", format="csv", content_base64="AAAAA")

    with pytest.raises(history.HistoricalImportError, match="file is too large"):
        history._validate_documents([document])

    assert calls == 0


def test_sec_d2_oversized_encoded_aggregate_rejects_before_any_decode(monkeypatch) -> None:
    calls = 0

    def unexpected_decode(*args, **kwargs):
        nonlocal calls
        calls += 1
        raise AssertionError("b64decode must not run")

    monkeypatch.setattr(history, "MAX_DOCUMENT_BASE64_CHARS", 100)
    monkeypatch.setattr(history, "MAX_UPLOAD_BASE64_CHARS", 6)
    monkeypatch.setattr(history.base64, "b64decode", unexpected_decode)
    documents = [
        _constructed_document(name="patients.csv", format="csv", content_base64="AAAA"),
        _constructed_document(name="appointments.csv", format="csv", content_base64="AAAA"),
    ]

    with pytest.raises(history.HistoricalImportError, match="too large in total"):
        history._validate_documents(documents)

    assert calls == 0


def test_sec_d3_invalid_base64_within_bound_reaches_decoder(monkeypatch) -> None:
    original_decode = base64.b64decode
    calls = 0

    def counting_decode(*args, **kwargs):
        nonlocal calls
        calls += 1
        return original_decode(*args, **kwargs)

    monkeypatch.setattr(history.base64, "b64decode", counting_decode)
    document = _constructed_document(name="patients.csv", format="csv", content_base64="%%%")

    with pytest.raises(history.HistoricalImportError, match="invalid file encoding"):
        history._validate_documents([document])

    assert calls == 1


def test_sec_d4_actual_decoded_size_check_remains_authoritative(monkeypatch) -> None:
    monkeypatch.setattr(history, "MAX_DOCUMENT_BYTES", 4)
    monkeypatch.setattr(history, "MAX_UPLOAD_BYTES", 100)
    monkeypatch.setattr(history, "MAX_DOCUMENT_BASE64_CHARS", 8)
    monkeypatch.setattr(history, "MAX_UPLOAD_BASE64_CHARS", 100)
    document = _constructed_document(
        name="patients.csv",
        format="csv",
        content_base64=base64.b64encode(b"ABCDEF").decode("ascii"),
    )

    assert len(document.content_base64) == 8
    with pytest.raises(history.HistoricalImportError, match="file is too large"):
        history._validate_documents([document])


def test_sec_d5_small_csv_and_xlsx_continue_to_validate_and_parse() -> None:
    csv_raw = b"full_name,phone\nAlice,01000000000\n"
    csv_document = HistoricalImportDocument(
        name="patients.csv",
        format="csv",
        content_base64=base64.b64encode(csv_raw).decode("ascii"),
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "patients"
    sheet.append(["full_name", "phone"])
    sheet.append(["Bob", "01100000000"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    xlsx_document = HistoricalImportDocument(
        name="history.xlsx",
        format="xlsx",
        content_base64=base64.b64encode(buffer.getvalue()).decode("ascii"),
    )

    history._validate_documents([csv_document, xlsx_document])
    rows = list(history._iter_rows([csv_document, xlsx_document]))

    assert any(source_file == "patients.csv" and sheet_name == "patients" for source_file, sheet_name, _, _ in rows)
    assert any(source_file == "history.xlsx" and sheet_name == "patients" for source_file, sheet_name, _, _ in rows)
